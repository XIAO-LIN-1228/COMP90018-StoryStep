from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "assignment1_project_plan_draft.md"
OUT_DIR = ROOT / "output" / "docx"
OUT_PATH = OUT_DIR / "StoryStep_Assignment1_Project_Plan.docx"

NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "222222"
MUTED = "5B6573"
LIGHT_BLUE = "EAF1F8"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"
GRID = "CBD3DD"


def clean(text: str) -> str:
    replacements = {
        "–": "-", "—": "-", "‑": "-", "’": "'", "“": '"', "”": '"',
        "…": "...", "→": "to", "↗": "NE", "`": "",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


def set_font(run, size=None, bold=None, italic=None, color=INK, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_in, indent_dxa=120):
    widths_dxa = [round(w * 1440) for w in widths_in]
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_in[idx])
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_columns(section, number=1, space=360):
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")
    if cols:
        cols_el = cols[0]
    else:
        cols_el = OxmlElement("w:cols")
        sect_pr.append(cols_el)
    cols_el.set(qn("w:num"), str(number))
    cols_el.set(qn("w:space"), str(space))


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, separate, text, end):
        run._r.append(el)
    set_font(run, size=9, color=MUTED)


def configure_section(section, first_page=False):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    section.different_first_page_header_footer = first_page
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.text = "STORYSTEP  |  COMP90018 PROJECT PLAN"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    for run in hp.runs:
        set_font(run, size=8.5, bold=True, color=MUTED)
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    add_page_field(fp)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    specs = {
        "Heading 1": (16, BLUE, 14, 6),
        "Heading 2": (12.5, BLUE, 10, 4),
        "Heading 3": (11, DARK_BLUE, 7, 3),
    }
    for name, (size, color, before, after) in specs.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(10.25)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.08


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(56)
    p.paragraph_format.space_after = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("COMP90018")
    set_font(r, size=12, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Assignment 1 - Project Plan")
    set_font(r, size=28, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(30)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("STORYSTEP")
    set_font(r, size=19, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Every step reveals a story.")
    set_font(r, size=12, italic=True, color=MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(30)
    r = p.add_run("Context-aware Android tourism companion")
    set_font(r, size=11, color=MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Group [INSERT GROUP NUMBER]")
    set_font(r, size=13, bold=True, color=NAVY)

    table = doc.add_table(rows=6, cols=3)
    table.style = "Table Grid"
    headers = ["Group member", "Student number", "University email"]
    for i, value in enumerate(headers):
        table.cell(0, i).text = value
    for row_idx in range(1, 6):
        values = [f"[Member {row_idx} full name]", "[Student number]", "[Email]"]
        for col_idx, value in enumerate(values):
            table.cell(row_idx, col_idx).text = value
    set_table_geometry(table, [2.15, 1.45, 2.9])
    repeat_table_header(table.rows[0])
    for row_idx, row in enumerate(table.rows):
        prevent_row_split(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_idx == 0:
                shade_cell(cell, NAVY)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.0
                for run in para.runs:
                    set_font(run, size=9.2, bold=row_idx == 0, color=WHITE if row_idx == 0 else INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("The University of Melbourne")
    set_font(r, size=10.5, bold=True, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Due 1 September 2026")
    set_font(r, size=10, color=MUTED)


CONTRIBUTIONS = [
    ("[Member 1] - Coordination, architecture and location", [
        "Maintain the delivery board, milestone checklist, meeting decisions and risk register.",
        "Lead package-name correction and establish the shared architecture.",
        "Implement location permission handling and the foreground exploration location source.",
        "Implement and test distance, route accumulation and arrival detection.",
        "Integrate the exploration ViewModel and review session persistence.",
    ]),
    ("[Member 2] - User experience and attraction discovery", [
        "Produce user flows, wireframes and accessible XML screen designs.",
        "Implement attraction lists, details, loading, empty and error states.",
        "Implement selection, saving feedback and clear session controls.",
        "Test screen sizes, font scales, insets and Material consistency.",
        "Review recommendation presentation and remote error handling.",
    ]),
    ("[Member 3] - Motion sensing and directional guidance", [
        "Implement lifecycle-safe accelerometer and magnetometer sources.",
        "Develop testable movement-state and heading calculations.",
        "Implement bearing, relative direction and compass behaviour.",
        "Add smoothing, accuracy feedback and missing-sensor fallbacks.",
        "Review distance calculations and lead physical-device testing.",
    ]),
    ("[Member 4] - Local data, sessions and history", [
        "Design Room entities and data access for attractions, saves and sessions.",
        "Implement session coordination, duration, route summary and visits.",
        "Implement history and saved-attraction screens and empty states.",
        "Test local persistence and session-state transitions.",
        "Review sensor lifecycle behaviour and support regression testing.",
    ]),
    ("[Member 5] - Remote data and contextual recommendations", [
        "Evaluate attraction and weather services for coverage and reliability.",
        "Implement Retrofit/OkHttp clients, mapping and repository integration.",
        "Implement deterministic ranking using context and preferences.",
        "Add caching and graceful offline, slow and invalid-response handling.",
        "Prepare ranking fixtures and coordinate technical documentation.",
    ]),
]


def add_contributions(doc):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section)
    set_columns(section, 2, 420)
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(0)
    p.add_run("2. Potential contribution breakdown")
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(5)
    r = intro.add_run("Primary ownership identifies who drives a deliverable; each major component also has a peer reviewer to share knowledge and reduce integration risk.")
    set_font(r, size=9, italic=True, color=MUTED)
    for idx, (title, bullets) in enumerate(CONTRIBUTIONS):
        if idx == 3:
            doc.add_paragraph().add_run().add_break(WD_BREAK.COLUMN)
        p = doc.add_paragraph(style="Heading 2")
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        set_font(r, size=9.4, bold=True, color=DARK_BLUE)
        for item in bullets:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(1.3)
            p.paragraph_format.line_spacing = 0.98
            for run in p.runs:
                set_font(run, size=8.55, color=INK)
            if not p.runs:
                r = p.add_run(item)
                set_font(r, size=8.55, color=INK)
            else:
                p.runs[0].text = item
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Shared responsibility")
    set_font(r, size=9.4, bold=True, color=DARK_BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 0.98
    r = p.add_run("All members review pull requests, test their work, join physical-device sessions, update documentation and contribute to the demonstration. Work is rebalanced at each milestone using deliverables and remaining estimates rather than commit counts.")
    set_font(r, size=8.55, color=INK)


def parse_markdown_blocks(text):
    lines = text.splitlines()
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [clean(c.strip()) for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[-: ]+", c or " ") for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
            continue
        if line.startswith("### "):
            blocks.append(("h3", clean(line[4:].strip())))
            i += 1
            continue
        if line.startswith("## "):
            blocks.append(("h2", clean(line[3:].strip())))
            i += 1
            continue
        if line.startswith("- "):
            blocks.append(("bullet", clean(line[2:].strip())))
            i += 1
            continue
        paragraph = [line.strip()]
        i += 1
        while i < len(lines):
            nxt = lines[i].rstrip()
            if not nxt.strip():
                i += 1
                break
            if nxt.startswith(("## ", "### ", "- ", "|")):
                break
            paragraph.append(nxt.strip())
            i += 1
        blocks.append(("p", clean(" ".join(paragraph))))
    return blocks


def table_widths(headers):
    joined = " ".join(headers).lower()
    if "target period" in joined:
        return [1.2, 1.05, 2.1, 2.15]
    if "likelihood" in joined:
        return [1.5, 1.0, 1.65, 2.35]
    if "assignment 1 criterion" in joined:
        return [1.35, 1.45, 2.45, 1.25]
    if len(headers) == 5:
        return [1.25, 1.55, 1.65, 1.25, 0.8]
    return [6.5 / len(headers)] * len(headers)


def add_table(doc, rows):
    if not rows:
        return
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    widths = table_widths(rows[0])
    for row_idx, values in enumerate(rows):
        for col_idx in range(cols):
            cell = table.cell(row_idx, col_idx)
            cell.text = values[col_idx] if col_idx < len(values) else ""
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_idx == 0:
                shade_cell(cell, NAVY)
            elif row_idx % 2 == 0:
                shade_cell(cell, LIGHT_GRAY)
            for para in cell.paragraphs:
                para.paragraph_format.space_before = Pt(0)
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.0
                for run in para.runs:
                    set_font(run, size=8.1, bold=row_idx == 0, color=WHITE if row_idx == 0 else INK)
        prevent_row_split(table.rows[row_idx])
    repeat_table_header(table.rows[0])
    set_table_geometry(table, widths)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)


def add_blocks(doc, blocks):
    for kind, value in blocks:
        if kind == "h2":
            p = doc.add_paragraph(style="Heading 1")
            p.add_run(value)
        elif kind == "h3":
            p = doc.add_paragraph(style="Heading 2")
            p.add_run(value)
        elif kind == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(value)
        elif kind == "table":
            add_table(doc, value)
        else:
            italic = value.startswith("*") and value.endswith("*")
            text = value.strip("*") if italic else value
            p = doc.add_paragraph()
            p.paragraph_format.keep_together = False
            r = p.add_run(text)
            if italic:
                set_font(r, size=9.5, italic=True, color=MUTED)


def extract_section(source, start_heading, end_heading=None):
    start = source.index(start_heading)
    end = source.index(end_heading, start) if end_heading else len(source)
    return source[start:end].strip()


def build():
    source = SOURCE.read_text(encoding="utf-8")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0], first_page=True)
    add_cover(doc)

    overview_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(overview_section)
    set_columns(overview_section, 1)
    overview = extract_section(source, "## 1. Project overview", "## 2. Potential contribution breakdown")
    add_blocks(doc, parse_markdown_blocks(overview))

    p = doc.add_paragraph(style="Heading 2")
    p.add_run("Project at a glance")
    table = doc.add_table(rows=6, cols=2)
    table.style = "Table Grid"
    items = [
        ("Platform", "Android; Java application code; XML layouts; minSdk 26"),
        ("Primary context", "Location, movement, orientation, weather, time and visit preferences"),
        ("Core experience", "Discover, select, explore, receive approximate guidance, unlock a story and retain history"),
        ("Data", "Room local persistence plus remote weather and attraction services"),
        ("Quality focus", "Testable calculations, foreground-only tracking, privacy and graceful degradation"),
    ]
    table.cell(0, 0).text = "Area"
    table.cell(0, 1).text = "Planned approach"
    repeat_table_header(table.rows[0])
    for cell in table.rows[0].cells:
        shade_cell(cell, NAVY)
        for para in cell.paragraphs:
            para.paragraph_format.space_after = Pt(0)
            for run in para.runs:
                set_font(run, size=9.2, bold=True, color=WHITE)
    for row_idx, (label, detail) in enumerate(items, start=1):
        table.cell(row_idx, 0).text = label
        table.cell(row_idx, 1).text = detail
        shade_cell(table.cell(row_idx, 0), LIGHT_BLUE)
        for col_idx, cell in enumerate(table.rows[row_idx].cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                for run in para.runs:
                    set_font(run, size=9.2, bold=col_idx == 0, color=NAVY if col_idx == 0 else INK)
    set_table_geometry(table, [1.35, 5.15])

    add_contributions(doc)

    detail_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(detail_section)
    set_columns(detail_section, 1)
    detail = extract_section(source, "## 3. Implementation plan and rationale", "## 7. Use of tools and assistance")
    add_blocks(doc, parse_markdown_blocks(detail))

    ack = extract_section(source, "## 7. Use of tools and assistance")
    add_blocks(doc, parse_markdown_blocks(ack))

    doc.core_properties.title = "StoryStep - Assignment 1 Project Plan"
    doc.core_properties.subject = "COMP90018 Mobile Computing Systems Programming"
    doc.core_properties.author = "StoryStep Project Group"
    doc.core_properties.keywords = "StoryStep, COMP90018, project plan"
    doc.core_properties.comments = ""
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build()
